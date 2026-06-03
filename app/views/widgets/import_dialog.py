from PySide6.QtCore import Qt
from PySide6.QtWidgets import (
    QComboBox,
    QDialog,
    QFileDialog,
    QFormLayout,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QMessageBox,
    QProgressBar,
    QPushButton,
    QScrollArea,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)

from app.utils.importer import (
    CAMPOS_EDITAVEIS,
    Importador,
    ImportResult,
    detectar_mapeamento,
    parse_csv,
    parse_xlsx,
)
from app.views.styles.theme import (
    ESTILO_BOTAO_AVISO,
    ESTILO_BOTAO_ERRO,
    ESTILO_BOTAO_FECHAR,
    ESTILO_BOTAO_PRIMARIO,
    ESTILO_BOTAO_SUCESSO,
    ESTILO_DIALOG,
    ESTILO_INPUT,
    ESTILO_INPUT_READONLY,
    ESTILO_TABELA_SIMPLES,
    ESTILO_TITULO_PAGINA,
)


class ImportDialog(QDialog):
    def __init__(self, parent, entity: str, entity_nome: str, service, session):
        super().__init__(parent)
        self._entity = entity
        self._entity_nome = entity_nome
        self._service = service
        self._session = session
        self._caminho = ""
        self._cabecalho: list[str] = []
        self._dados: list[list[str]] = []
        self._mapa_colunas: list[str] = []
        self._combo_mapas: list[QComboBox] = []

        self.setWindowTitle(f"Importar {entity_nome}")
        self.setMinimumSize(800, 720)
        self.setStyleSheet(ESTILO_DIALOG)

        self._setup_ui()

    def _setup_ui(self):
        layout = QVBoxLayout(self)
        layout.setSpacing(12)
        layout.setContentsMargins(20, 20, 20, 20)

        titulo = QLabel(f"📥 Importar {self._entity_nome}")
        titulo.setStyleSheet(ESTILO_TITULO_PAGINA)
        layout.addWidget(titulo)

        # File selection
        file_layout = QHBoxLayout()
        self._lbl_arquivo = QLineEdit()
        self._lbl_arquivo.setReadOnly(True)
        self._lbl_arquivo.setPlaceholderText("Selecione um arquivo CSV ou XLSX...")
        self._lbl_arquivo.setStyleSheet(ESTILO_INPUT_READONLY)
        file_layout.addWidget(self._lbl_arquivo)

        btn_selecionar = QPushButton("📂 Selecionar")
        btn_selecionar.setStyleSheet(ESTILO_BOTAO_PRIMARIO)
        btn_selecionar.clicked.connect(self._selecionar_arquivo)
        file_layout.addWidget(btn_selecionar)
        layout.addLayout(file_layout)

        # Options (encoding, delimiter) - shown conditionally
        opts_layout = QHBoxLayout()
        opts_layout.setSpacing(8)

        self._lbl_encoding = QLabel("Codificação:")
        self._lbl_encoding.setStyleSheet("color: #c8c8d8; font-size: 12px; background: transparent;")
        self._combo_encoding = QComboBox()
        self._combo_encoding.addItems(["UTF-8", "ISO-8859-1", "Windows-1252", "UTF-16"])
        self._combo_encoding.setStyleSheet(self._estilo_combo_opt())
        opts_layout.addWidget(self._lbl_encoding)
        opts_layout.addWidget(self._combo_encoding)

        self._lbl_delim = QLabel("Delimitador:")
        self._lbl_delim.setStyleSheet("color: #c8c8d8; font-size: 12px; background: transparent;")
        self._combo_delim = QComboBox()
        self._combo_delim.addItems([";", ",", "|", "\t"])
        self._combo_delim.setStyleSheet(self._estilo_combo_opt())
        opts_layout.addWidget(self._lbl_delim)
        opts_layout.addWidget(self._combo_delim)

        self._lbl_sheet = QLabel("Planilha:")
        self._lbl_sheet.setStyleSheet("color: #c8c8d8; font-size: 12px; background: transparent;")
        self._combo_sheet = QComboBox()
        self._combo_sheet.setStyleSheet(self._estilo_combo_opt())
        opts_layout.addWidget(self._lbl_sheet)
        opts_layout.addWidget(self._combo_sheet)

        opts_layout.addStretch()
        btn_analisar = QPushButton("🔍 Analisar")
        btn_analisar.setStyleSheet(ESTILO_BOTAO_AVISO)
        btn_analisar.clicked.connect(self._analisar)
        opts_layout.addWidget(btn_analisar)
        layout.addLayout(opts_layout)

        # Preview header
        preview_header = QHBoxLayout()
        lbl_preview = QLabel("Pré-visualização:")
        lbl_preview.setStyleSheet("color: #94949f; font-size: 12px; background: transparent;")
        preview_header.addWidget(lbl_preview)
        preview_header.addStretch()
        self._lbl_total_linhas = QLabel("")
        self._lbl_total_linhas.setStyleSheet("color: #475569; font-size: 12px; background: transparent;")
        preview_header.addWidget(self._lbl_total_linhas)
        layout.addLayout(preview_header)

        # Preview table
        self._preview_table = QTableWidget()
        self._preview_table.setStyleSheet(ESTILO_TABELA_SIMPLES)
        self._preview_table.setEditTriggers(QTableWidget.NoEditTriggers)
        self._preview_table.setAlternatingRowColors(True)
        self._preview_table.verticalHeader().setVisible(False)
        self._preview_table.setMaximumHeight(450)
        layout.addWidget(self._preview_table)

        # Column mapping
        lbl_map = QLabel("Mapeamento de Colunas:")
        lbl_map.setStyleSheet("color: #94949f; font-size: 12px; background: transparent;")
        layout.addWidget(lbl_map)

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setMaximumHeight(140)
        scroll.setStyleSheet("QScrollArea { border: none; background: transparent; }")

        self._map_widget = QWidget()
        self._map_layout = QVBoxLayout(self._map_widget)
        self._map_layout.setContentsMargins(0, 0, 0, 0)
        self._map_layout.setSpacing(4)
        scroll.setWidget(self._map_widget)
        layout.addWidget(scroll)

        # Progress
        self._progress = QProgressBar()
        self._progress.setVisible(False)
        self._progress.setStyleSheet("""
            QProgressBar { background-color: #1e1e2e; border: 1px solid #2a2a3e;
                           border-radius: 6px; text-align: center; color: #e8e8f0;
                           font-size: 11px; height: 20px; }
            QProgressBar::chunk { background-color: #6366f1; border-radius: 5px; }
        """)
        layout.addWidget(self._progress)

        # Status
        self._lbl_status = QLabel("")
        self._lbl_status.setStyleSheet("color: #475569; font-size: 12px; background: transparent;")
        layout.addWidget(self._lbl_status)

        # Buttons
        btn_layout = QHBoxLayout()
        btn_layout.addStretch()

        self._btn_importar = QPushButton("📥 Importar")
        self._btn_importar.setStyleSheet(ESTILO_BOTAO_SUCESSO)
        self._btn_importar.setEnabled(False)
        self._btn_importar.clicked.connect(self._importar)
        btn_layout.addWidget(self._btn_importar)

        btn_fechar = QPushButton("Fechar")
        btn_fechar.setStyleSheet(ESTILO_BOTAO_FECHAR)
        btn_fechar.clicked.connect(self.reject)
        btn_layout.addWidget(btn_fechar)

        layout.addLayout(btn_layout)

        # Initial visibility
        self._combo_sheet.setVisible(False)
        self._lbl_sheet.setVisible(False)

    def _estilo_combo_opt(self):
        return (
            "QComboBox { background-color: #1e1e2e; color: #e8e8f0;"
            " border: 1px solid #2a2a3e; border-radius: 6px;"
            " padding: 4px 8px; font-size: 11px; min-height: 20px; }"
        )

    def _selecionar_arquivo(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "Selecionar Arquivo", "",
            "Arquivos Suportados (*.csv *.xlsx);;CSV (*.csv);;Excel (*.xlsx);;Todos (*.*)",
        )
        if not path:
            return
        self._caminho = path
        self._lbl_arquivo.setText(path)
        is_csv = path.lower().endswith(".csv")
        self._lbl_encoding.setVisible(is_csv)
        self._combo_encoding.setVisible(is_csv)
        self._lbl_delim.setVisible(is_csv)
        self._combo_delim.setVisible(is_csv)
        self._lbl_sheet.setVisible(not is_csv)
        self._combo_sheet.setVisible(not is_csv)
        self._combo_sheet.clear()
        self._btn_importar.setEnabled(False)
        self._preview_table.setRowCount(0)
        self._preview_table.setColumnCount(0)

        if not is_csv:
            try:
                from openpyxl import load_workbook
                wb = load_workbook(path, read_only=True)
                self._combo_sheet.addItems(wb.sheetnames)
                wb.close()
            except Exception as e:
                QMessageBox.warning(self, "Aviso", f"Não foi possível ler o arquivo:\n{e}")

    def _analisar(self):
        if not self._caminho:
            QMessageBox.warning(self, "Aviso", "Selecione um arquivo primeiro.")
            return

        try:
            is_csv = self._caminho.lower().endswith(".csv")
            if is_csv:
                encoding = self._combo_encoding.currentText()
                delim = self._combo_delim.currentText()
                if delim == "\\t":
                    delim = "\t"
                cabecalho, dados = parse_csv(self._caminho, encoding, delim)
            else:
                sheet = self._combo_sheet.currentText() or None
                cabecalho, dados = parse_xlsx(self._caminho, sheet)

            if not cabecalho:
                QMessageBox.warning(self, "Aviso", "Arquivo vazio ou cabeçalho não encontrado.")
                return

            self._cabecalho = cabecalho
            self._dados = dados
            total = len(dados)
            self._lbl_total_linhas.setText(f"{total} linha(s) de dados")

            # Auto-detect mapping
            mapa = detectar_mapeamento(cabecalho, self._entity)
            self._mapa_colunas = [mapa.get(i) for i in range(len(cabecalho))]

            # Preview
            self._preview_table.setColumnCount(len(cabecalho))
            self._preview_table.setHorizontalHeaderLabels(cabecalho)
            n_preview = len(dados)
            self._preview_table.setRowCount(n_preview)
            for i in range(n_preview):
                for j in range(len(cabecalho)):
                    val = dados[i][j] if j < len(dados[i]) else ""
                    item = QTableWidgetItem(val)
                    item.setForeground(Qt.gray if not val else Qt.white)
                    self._preview_table.setItem(i, j, item)
            self._preview_table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)

            # Build mapping combos
            self._construir_mapeamento(cabecalho)

            self._btn_importar.setEnabled(True)
            self._lbl_status.setText("✅ Arquivo analisado. Verifique o mapeamento e clique em Importar.")
            self._lbl_status.setStyleSheet("color: #34d399; font-size: 12px; background: transparent;")

        except Exception as e:
            QMessageBox.critical(self, "Erro", f"Erro ao analisar arquivo:\n{e}")
            self._lbl_status.setText(f"❌ Erro: {e}")
            self._lbl_status.setStyleSheet("color: #f87171; font-size: 12px; background: transparent;")

    def _construir_mapeamento(self, cabecalho: list[str]):
        # Clear old mapping widgets
        while self._map_layout.count():
            item = self._map_layout.takeAt(0)
            if item.widget():
                item.widget().deleteLater()
        self._combo_mapas.clear()

        campos = CAMPOS_EDITAVEIS.get(self._entity, [])
        opcoes = ["─ Ignorar ─"] + campos

        for i, nome_col in enumerate(cabecalho):
            row = QHBoxLayout()
            row.setSpacing(8)

            label = QLabel(f"{nome_col}  →")
            label.setStyleSheet("color: #c8c8d8; font-size: 12px; background: transparent; min-width: 120px;")
            row.addWidget(label)

            combo = QComboBox()
            combo.addItems(opcoes)
            combo.setStyleSheet(
                "QComboBox { background-color: #1e1e2e; color: #e8e8f0;"
                " border: 1px solid #2a2a3e; border-radius: 6px;"
                " padding: 4px 8px; font-size: 11px; min-height: 22px; }"
            )
            mapped = self._mapa_colunas[i] if i < len(self._mapa_colunas) else None
            if mapped and mapped in campos:
                idx = opcoes.index(mapped)
                combo.setCurrentIndex(idx)
            combo.currentIndexChanged.connect(self._atualizar_mapa)
            row.addWidget(combo)
            self._combo_mapas.append(combo)

            self._map_layout.addLayout(row)

    def _atualizar_mapa(self):
        for i, combo in enumerate(self._combo_mapas):
            texto = combo.currentText()
            if texto == "─ Ignorar ─":
                self._mapa_colunas[i] = None
            else:
                self._mapa_colunas[i] = texto

    def _importar(self):
        if not self._dados:
            return

        self._btn_importar.setEnabled(False)
        self._progress.setVisible(True)
        self._progress.setMaximum(0)
        self._progress.setValue(0)
        self._lbl_status.setText("Importando...")
        self._lbl_status.setStyleSheet("color: #fbbf24; font-size: 12px; background: transparent;")

        importador = Importador(self._entity, self._service, self._session)
        resultado = importador.importar_linhas(
            self._dados, self._mapa_colunas, self._cabecalho
        )

        self._progress.setMaximum(100)
        self._progress.setValue(100)

        # Show result
        if resultado.erros:
            msg_erros = "\n".join(
                f"  Linha {num}: {err}" for num, err in resultado.erros[:20]
            )
            if len(resultado.erros) > 20:
                msg_erros += f"\n  ... e mais {len(resultado.erros) - 20} erro(s)"
            msg = (
                f"✅ {resultado.importados} de {resultado.total} registros importados.\n\n"
                f"❌ {len(resultado.erros)} erro(s):\n{msg_erros}"
            )
            self._lbl_status.setText(f"⚠️ {resultado.importados} importados, {len(resultado.erros)} erro(s)")
            self._lbl_status.setStyleSheet("color: #fbbf24; font-size: 12px; background: transparent;")
            QMessageBox.warning(self, "Resultado da Importação", msg)
        else:
            self._lbl_status.setText(f"✅ {resultado.importados} registros importados com sucesso!")
            self._lbl_status.setStyleSheet("color: #34d399; font-size: 12px; background: transparent;")
            QMessageBox.information(
                self, "Importação Concluída",
                f"{resultado.importados} registro(s) importado(s) com sucesso!"
            )

        self._progress.setVisible(False)
        self._btn_importar.setEnabled(True)

        if resultado.importados > 0:
            self.accept()
