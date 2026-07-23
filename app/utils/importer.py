import csv
import io
import os
import re
from dataclasses import dataclass, field
from typing import Any, Optional

from app.utils.sanitize import sanitizar


@dataclass
class ImportResult:
    total: int = 0
    importados: int = 0
    erros: list[tuple[int, str]] = field(default_factory=list)


def detectar_encoding(caminho: str) -> str:
    with open(caminho, "rb") as f:
        raw = f.read(4096)
    try:
        import chardet
        return chardet.detect(raw).get("encoding", "utf-8") or "utf-8"
    except ImportError:
        if raw.startswith(b"\xff\xfe"):
            return "utf-16-le"
        if raw.startswith(b"\xfe\xff"):
            return "utf-16-be"
        return "utf-8"


def _limpar_cabecalho(nome: str) -> str:
    nome = nome.strip().lower()
    nome = re.sub(r"[^a-z0-9\u00e0-\u00ff]", " ", nome)
    nome = re.sub(r"\s+", "_", nome).strip("_")
    return nome


COLUNAS_PADRAO: dict[str, dict[str, str]] = {
    "printers": {
        "patrimonio": "patrimonio",
        "patrimônio": "patrimonio",
        "n_patrimonio": "patrimonio",
        "n_patrimônio": "patrimonio",
        "num_patrimonio": "patrimonio",
        "modelo": "modelo",
        "marca": "marca",
        "serial": "serial",
        "numero_serial": "serial",
        "número_serial": "serial",
        "n_serial": "serial",
        "no_serial": "serial",
        "status": "status",
        "local": "local_atual",
        "local_atual": "local_atual",
        "ip": "ip_rede",
        "ip_rede": "ip_rede",
        "tecnico": "tecnico",
        "técnico": "tecnico",
        "tipo": "tipo",
        "observacao": "observacao",
        "observação": "observacao",
        "obs": "observacao",
        "mac": "mac_address",
        "mac_address": "mac_address",
    },
    "parts": {
        "codigo": "codigo",
        "código": "codigo",
        "cod": "codigo",
        "nome": "nome",
        "descricao": "descricao",
        "descrição": "descricao",
        "desc": "descricao",
        "quantidade": "quantidade_estoque",
        "qtd": "quantidade_estoque",
        "qty": "quantidade_estoque",
        "estoque": "quantidade_estoque",
        "estoque_minimo": "estoque_minimo",
        "est_min": "estoque_minimo",
        "minimo": "estoque_minimo",
        "mínimo": "estoque_minimo",
        "preco": "preco_unitario",
        "preço": "preco_unitario",
        "preco_unitario": "preco_unitario",
        "modelo_compativel": "modelo_compativel",
        "modelo": "modelo_compativel",
    },
    "companies": {
        "nome": "nome",
        "empresa": "nome",
        "razao_social": "nome",
        "razão_social": "nome",
        "cnpj": "cnpj",
        "cpf": "cnpj",
        "endereco": "endereco",
        "endereço": "endereco",
        "cidade": "cidade",
        "uf": "uf",
        "estado": "uf",
        "telefone": "telefone",
        "tel": "telefone",
        "fone": "telefone",
        "email": "email",
        "e_mail": "email",
        "tipo": "tipo",
        "observacao": "observacao",
        "observação": "observacao",
    },
}

CAMPOS_OBRIGATORIOS: dict[str, list[str]] = {
    "printers": ["patrimonio"],
    "parts": ["nome"],
    "companies": ["nome"],
}

CAMPOS_EDITAVEIS: dict[str, list[str]] = {
    "printers": ["patrimonio", "modelo", "marca", "serial", "status", "local_atual",
                 "ip_rede", "mac_address", "tecnico", "tipo", "observacao"],
    "parts": ["codigo", "nome", "descricao", "quantidade_estoque", "estoque_minimo",
              "preco_unitario", "modelo_compativel"],
    "companies": ["nome", "cnpj", "endereco", "cidade", "uf", "telefone", "email", "tipo", "observacao"],
}


def _valor_convertido(valor: str, campo: str) -> Any:
    if valor is None:
        return None
    valor = valor.strip()
    if not valor:
        return None
    if campo in ("quantidade_estoque", "estoque_minimo"):
        try:
            return int(float(valor))
        except ValueError:
            return 0
    if campo == "preco_unitario":
        valor = valor.replace(",", ".")
        try:
            return float(valor)
        except ValueError:
            return 0.0
    return valor


def parse_csv(caminho: str, encoding: str = "utf-8", delimiter: str = ";") -> tuple[list[str], list[list[str]]]:
    with open(caminho, encoding=encoding, errors="replace") as f:
        reader = csv.reader(f, delimiter=delimiter)
        linhas = list(reader)
    if not linhas:
        return [], []
    cabecalho = [c.lstrip("\ufeff") for c in linhas[0]]
    dados = linhas[1:]
    return cabecalho, dados


def parse_xlsx(caminho: str, sheet_name: Optional[str] = None) -> tuple[list[str], list[list[str]]]:
    from openpyxl import load_workbook
    wb = load_workbook(caminho, read_only=True, data_only=True)
    if sheet_name:
        ws = wb[sheet_name]
    else:
        ws = wb.active
    linhas = list(ws.iter_rows(values_only=True))
    if not linhas:
        return [], []
    cabecalho = [str(c or "") for c in linhas[0]]
    dados = []
    for row in linhas[1:]:
        dados.append([str(c or "") for c in row])
    return cabecalho, dados


def detectar_mapeamento(cabecalho: list[str], entity: str) -> dict[int, Optional[str]]:
    mapa: dict[int, Optional[str]] = {}
    colunas = COLUNAS_PADRAO.get(entity, {})
    chaves_ordenadas = sorted(colunas.keys(), key=len, reverse=True)
    usados: set[str] = set()
    for i, nome in enumerate(cabecalho):
        chave = _limpar_cabecalho(nome)
        campo = colunas.get(chave)
        if campo:
            mapa[i] = campo
            usados.add(campo)
        else:
            melhor_campo: Optional[str] = None
            for k in chaves_ordenadas:
                if k in chave and colunas[k] not in usados:
                    melhor_campo = colunas[k]
                    break
            mapa[i] = melhor_campo
            if melhor_campo:
                usados.add(melhor_campo)
    return mapa


def mapeamento_para_lista(mapa: dict[int, Optional[str]], entity: str) -> list[tuple[int, str, Optional[str]]]:
    result = []
    for idx in sorted(mapa.keys()):
        result.append((idx, f"Coluna {idx + 1}", mapa.get(idx)))
    return result


class Importador:
    def __init__(self, entity: str, service, session):
        self.entity = entity
        self.service = service
        self.session = session

    def importar_linhas(self, dados: list[list[str]], mapa_colunas: list[Optional[str]],
                        cabecalho: list[str]) -> ImportResult:
        result = ImportResult()
        result.total = len(dados)
        obrigatorios = CAMPOS_OBRIGATORIOS.get(self.entity, [])
        editaveis = CAMPOS_EDITAVEIS.get(self.entity, [])

        for idx_linha, linha in enumerate(dados):
            num_linha = idx_linha + 2
            dados_linha: dict[str, Any] = {}
            pulada = False
            for col_idx, campo in enumerate(mapa_colunas):
                if campo and col_idx < len(linha):
                    valor = _valor_convertido(linha[col_idx], campo)
                    if valor is not None:
                        dados_linha[campo] = valor

            for req in obrigatorios:
                if req not in dados_linha or dados_linha[req] in (None, ""):
                    result.erros.append((num_linha, f"Campo obrigatório '{req}' vazio"))
                    pulada = True
                    break
            if pulada:
                continue

            # Remove campos não editáveis
            dados_filtrados = {k: v for k, v in dados_linha.items() if k in editaveis}

            # Sanitiza strings
            for k, v in list(dados_filtrados.items()):
                if isinstance(v, str):
                    dados_filtrados[k] = sanitizar(v, self.entity.capitalize(), k)

            try:
                self._importar_um(dados_filtrados)
                result.importados += 1
            except Exception as e:
                result.erros.append((num_linha, str(e)))

        return result

    def _importar_um(self, dados: dict[str, Any]):
        if self.entity == "printers":
            patrimonio = dados.get("patrimonio")
            if not patrimonio:
                raise ValueError("patrimonio é obrigatório")
            existente = self.service.buscar_por_patrimonio(patrimonio)
            if existente:
                self.service.atualizar(existente, **dados)
            else:
                self.service.criar(**dados)
        elif self.entity == "parts":
            nome = dados.get("nome", "")
            if not nome:
                raise ValueError("nome é obrigatório")
            codigo = dados.pop("codigo", None) or self.service.gerar_codigo()
            existente = self.service.buscar_por_nome(nome)
            if existente:
                self.service.atualizar(existente, **dados)
            else:
                self.service.criar(
                    codigo=codigo,
                    nome=dados.pop("nome", nome),
                    descricao=dados.pop("descricao", ""),
                    modelo_compativel=dados.pop("modelo_compativel", ""),
                    quantidade=dados.pop("quantidade_estoque", 0),
                    estoque_minimo=dados.pop("estoque_minimo", 1),
                )
        elif self.entity == "companies":
            nome = dados.get("nome", "")
            if not nome:
                raise ValueError("nome é obrigatório")
            existente = self.service.buscar_por_nome(nome)
            if existente:
                self.service.atualizar(existente, **dados)
            else:
                self.service.criar(
                    nome=nome,
                    cnpj=dados.pop("cnpj", ""),
                    telefone=dados.pop("telefone", ""),
                    email=dados.pop("email", ""),
                    tipo=dados.pop("tipo", "Cliente"),
                )
        else:
            raise ValueError(f"Entidade desconhecida: {self.entity}")

    def preview_data(self, cabecalho: list[str], dados: list[list[str]],
                     mapa_colunas: list[Optional[str]]) -> tuple[list[str], list[list[str]]]:
        colunas_exibidas = []
        indices_exibidos = []
        for i, campo in enumerate(mapa_colunas):
            if campo:
                colunas_exibidas.append(campo)
                indices_exibidos.append(i)

        if not colunas_exibidas:
            return [], []

        rows = []
        for linha in dados[:20]:
            row = []
            for idx in indices_exibidos:
                val = linha[idx] if idx < len(linha) else ""
                row.append(val)
            rows.append(row)

        return colunas_exibidas, rows
